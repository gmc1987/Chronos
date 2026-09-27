from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


OUTPUT = Path('/Users/gaomingchao/git/Chronos/docs/Chronos教育行业全业务展示与操作手册.docx')
BLUE = '1F4E78'
LIGHT_BLUE = 'D9EAF7'
PALE = 'F3F6F9'
GREEN = '2E7D32'
ORANGE = 'C65911'
GRAY = '666666'


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        tc_pr.append(shd)
    shd.set(qn('w:fill'), fill)


def set_cell_text(cell, text, bold=False, color=None, size=9):
    cell.text = ''
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(str(text))
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = 'Heiti SC'
    run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Heiti SC')
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = 'Table Grid'
    for idx, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[idx], header, True, 'FFFFFF', 9)
        set_cell_shading(table.rows[0].cells[idx], BLUE)
    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            set_cell_text(cells[idx], value, False, None, 8.5)
            if row_index % 2 == 1:
                set_cell_shading(cells[idx], PALE)
    if widths:
        for row in table.rows:
            for idx, width in enumerate(widths):
                row.cells[idx].width = Cm(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.keep_with_next = True
    return p


def add_steps(doc, steps):
    for idx, step in enumerate(steps, 1):
        # 使用显式序号，避免 Word/LibreOffice 将不同业务章节合并成一条连续编号。
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.5)
        p.paragraph_format.first_line_indent = Cm(-0.5)
        p.add_run(f'{idx}. ')
        p.add_run(step)


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.left_indent = Cm(0.5)
        p.add_run(item)


def add_note(doc, title, text, color=LIGHT_BLUE):
    table = doc.add_table(rows=1, cols=1)
    table.style = 'Table Grid'
    set_cell_shading(table.cell(0, 0), color)
    p = table.cell(0, 0).paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(f'{title}：')
    r.bold = True
    r.font.name = 'Heiti SC'
    r._element.rPr.rFonts.set(qn('w:eastAsia'), 'Heiti SC')
    r = p.add_run(text)
    r.font.name = 'Heiti SC'
    r._element.rPr.rFonts.set(qn('w:eastAsia'), 'Heiti SC')
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_module(doc, title, role, entry, data, steps, expected, notes=None):
    add_heading(doc, title, 2)
    add_table(doc, ['项目', '说明'], [
        ('建议账号', role),
        ('入口', entry),
        ('展示数据', data),
    ], [3.2, 13.8])
    add_heading(doc, '操作步骤', 3)
    add_steps(doc, steps)
    add_heading(doc, '预期结果', 3)
    add_bullets(doc, expected)
    if notes:
        add_note(doc, '注意', notes, 'FFF2CC')


def set_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run('第 ')
    fld_char1 = OxmlElement('w:fldChar')
    fld_char1.set(qn('w:fldCharType'), 'begin')
    instr_text = OxmlElement('w:instrText')
    instr_text.set(qn('xml:space'), 'preserve')
    instr_text.text = ' PAGE '
    fld_char2 = OxmlElement('w:fldChar')
    fld_char2.set(qn('w:fldCharType'), 'end')
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)
    paragraph.add_run(' 页')


doc = Document()
section = doc.sections[0]
section.top_margin = Cm(2.2)
section.bottom_margin = Cm(2.0)
section.left_margin = Cm(2.2)
section.right_margin = Cm(2.0)

styles = doc.styles
normal = styles['Normal']
normal.font.name = 'Heiti SC'
normal._element.rPr.rFonts.set(qn('w:eastAsia'), 'Heiti SC')
normal.font.size = Pt(10.5)
normal.paragraph_format.line_spacing = 1.35
normal.paragraph_format.space_after = Pt(6)
for style_name, size, color in [('Title', 30, BLUE), ('Subtitle', 15, GRAY), ('Heading 1', 20, BLUE), ('Heading 2', 15, BLUE), ('Heading 3', 11, ORANGE)]:
    style = styles[style_name]
    style.font.name = 'Heiti SC'
    style._element.rPr.rFonts.set(qn('w:eastAsia'), 'Heiti SC')
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor.from_string(color)
    style.font.bold = True

header = section.header.paragraphs[0]
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
header_run = header.add_run('Chronos AI Native 智慧校园 · 全业务操作手册')
header_run.font.size = Pt(8)
header_run.font.color.rgb = RGBColor(120, 120, 120)
set_page_number(section.footer.paragraphs[0])

# Cover
doc.add_paragraph().paragraph_format.space_after = Pt(70)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run('CHRONOS')
r.bold = True
r.font.size = Pt(16)
r.font.color.rgb = RGBColor.from_string(BLUE)
p = doc.add_paragraph(style='Title')
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run('教育行业全业务展示与操作手册')
p = doc.add_paragraph(style='Subtitle')
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run('面向中专 / 技工院校的一体化业务演示与验收指南')
doc.add_paragraph().paragraph_format.space_after = Pt(55)
add_table(doc, ['文档属性', '内容'], [
    ('版本', 'V1.0（完整展示数据版）'),
    ('适用数据库', 'ChronosEducation / PostgreSQL'),
    ('演示学校', '南城现代职业技术学校'),
    ('数据脚本', 'education-complete-showcase-data.sql'),
    ('编制日期', '2026-09-27'),
    ('用途', '产品展示、全业务线人工验收、培训与售前演示'),
], [4, 12])
add_note(doc, '安全提示', '本文账号与统一密码只用于本地演示环境。严禁将展示账号、测试密码或模拟业务数据部署到生产环境。', 'FCE4D6')
doc.add_page_break()

add_heading(doc, '使用说明', 1)
doc.add_paragraph('本手册以当前仓库、ChronosEducation 实际表结构和已实现页面为基线，覆盖平台治理、门户、流程、消息、教务、排课、教学、作业、考试、成绩、会议、请假、家校、督导、数据中心、知识库与集成中心。建议按第 4 章给出的端到端顺序执行，既可用于完整演示，也可作为人工验收清单。')
add_note(doc, '边界说明', '模拟数据保证页面有可展示的主数据、过程数据和结果数据；需要真实第三方密钥的短信、企业微信、DeepSeek、外部集成凭据不写入脚本。集成连接器和运行记录可展示，凭据必须由管理员在受保护环境中配置。')

add_heading(doc, '目录', 1)
add_table(doc, ['章节', '主题'], [
    ('1', '环境准备与数据初始化'), ('2', '展示账号与角色分工'), ('3', '数据资产总览'),
    ('4', '推荐的完整业务演示路线'), ('5', '平台治理与权限'), ('6', '门户与通知公告'),
    ('7', '流程中心与表单'), ('8', '教务基础数据'), ('9', '走班排课'),
    ('10', '教学中心'), ('11', '作业中心'), ('12', '考试中心'), ('13', '成绩中心'),
    ('14', '会议中心'), ('15', '请假与教室预约'), ('16', '家校协同'),
    ('17', '教学督导与数据中心'), ('18', 'AI、知识库与集成中心'),
    ('19', '验收清单与问题定位'),
], [2.2, 13.8])
doc.add_page_break()

add_heading(doc, '1. 环境准备与数据初始化', 1)
add_heading(doc, '1.1 前置条件', 2)
add_bullets(doc, [
    'Docker 中 PostgreSQL 容器名称为 Chronos，数据库为 ChronosEducation。',
    '后端 education-app 与前端 Chronos-UI 已按教育行业模板启动。',
    '涉及附件上传时启动 MinIO；本地约定账号 admin，密码 minio123456。',
    '后端加密密钥 CHRONOS_ENCRYPTION_KEY 已通过环境变量或系统密钥管理器配置。',
])
add_heading(doc, '1.2 导入展示数据', 2)
add_steps(doc, [
    '确认当前连接的是 ChronosEducation，而不是医院模板的 Chronos 数据库。',
    '在仓库根目录执行：docker exec -i Chronos psql -v ON_ERROR_STOP=1 -U Chronos -d ChronosEducation < Chronos/scripts/education-complete-showcase-data.sql。',
    '看到最后输出 COMMIT 即表示导入成功；脚本可重复执行，不会重复生成同一批记录。',
    '启动 education-app 和前端，先用 showcase.admin 检查管理后台首页、菜单和各模块列表。',
])
add_note(doc, '禁止事项', '不要在生产库执行本脚本。脚本会创建已知密码的展示账号，并写入大量教学、考试、成绩、消息和流程案例。', 'FCE4D6')

add_heading(doc, '2. 展示账号与角色分工', 1)
doc.add_paragraph('五个统一账号密码均为 ChronosExam@2026。管理员入口为 /admin/login，门户入口为 /login。')
add_table(doc, ['账号', '身份', '主要用途', '关键角色'], [
    ('showcase.admin', '平台管理员', '系统配置、权限、流程、所有业务后台', 'SUPER_ADMIN'),
    ('showcase.teacher', '任课教师 / 班主任', '备课、作业、课表、请假、会议、流程', 'EDU_TEACHER / EDU_CLASS_ADVISOR / WORKFLOW_USER'),
    ('showcase.student', '学生', '个人课表、作业、考试、成绩、请假、流程', 'ROLE_PLATFORM_USER / WORKFLOW_USER'),
    ('showcase.parent', '家长', '家校通知、回执、学生信息与成绩查看', 'ROLE_PLATFORM_USER'),
    ('showcase.reviewer', '教务审核人', '教学审核、成绩审核发布、流程管理', 'EDU_ACADEMIC_APPROVER / EDU_GRADE_REVIEWER / WORKFLOW_ADMIN'),
], [3.3, 2.6, 5.6, 5.5])
add_note(doc, '权限验收原则', '超级管理员能看到所有内容，不代表普通用户权限正确。流程办理、学生门户和家长门户必须分别使用对应展示账号验证。')

add_heading(doc, '3. 数据资产总览', 1)
add_table(doc, ['数据域', '数量', '展示重点'], [
    ('组织单元 / 教职工', '8 / 22', '学校、校区、教务处、专业系部及岗位人员'),
    ('教师 / 学生 / 家长', '14 / 60 / 60', '完整身份档案和家长绑定'),
    ('学期 / 专业 / 行政班', '3 / 6 / 6', '当前学期与中职专业班级'),
    ('课程 / 科目 / 教学任务', '16 / 16 / 17', '含跨班合班公共课'),
    ('教室 / 课表', '13 / 21', '普通教室、实训室、不可用时段与冲突案例'),
    ('作业', '2', '发布、提交和批改链路'),
    ('考试计划 / 场次 / 考场', '3 / 4 / 3', '考试编排、监考、换班和分析'),
    ('成绩册', '1', '审核、发布、更正、补考与重修'),
    ('流程定义 / 实例 / 任务', '10 / 22 / 23', '发起、审批、退回、事故与恢复'),
    ('会议 / 通知公告', '1 / 2', '会议全周期与分范围发布'),
    ('知识文档 / 集成连接器', '1 / 2', '知识检索与同步运维展示'),
    ('数据质量问题', '2', '发现、指派、处理和复核'),
], [5.1, 2.2, 9.7])

add_heading(doc, '4. 推荐的完整业务演示路线', 1)
add_steps(doc, [
    '管理员建立组织、用户、角色、菜单和原子权限，确认字典与行业模板。',
    '教务人员建立学期、专业、班级、课程、师生档案和教学任务。',
    '排课人员配置节次、教室、教师禁排和规则，生成方案，检查冲突后发布课表。',
    '教师从教学任务进入教学计划、教案、集体备课、课件、题库、知识点和教研活动。',
    '教师发布作业，学生提交，教师批改并发布结果。',
    '考试管理员创建考试计划、场次和考场，安排监考；教师可发起监考调换。',
    '考试完成后录入逐题得分，成绩中心汇总成绩册，审核人审核并发布。',
    '班主任发布班级通知，家长查看并回执；学生提交请假和教室预约流程。',
    '组织会议并完成会议纪要、材料和行动项；督导执行听课与整改闭环。',
    '数据中心查看指标、质量问题和报表任务；集成中心查看同步运行和死信。',
])

add_heading(doc, '5. 平台治理与权限', 1)
add_module(doc, '5.1 用户、角色与组织', 'showcase.admin', '管理后台 → 用户管理 / 角色管理 / 组织架构 / 通讯录', '南城现代职业技术学校、8 个组织单元、22 名教职工、五类展示账号', [
    '在组织架构中展开学校与校区，确认教务处、专业系部层级。',
    '进入用户管理搜索 showcase.，逐个打开详情，确认状态、员工和角色回显。',
    '进入角色管理，分别查看教师、班主任、教务审核和流程角色的授权。',
    '进入通讯录按组织筛选人员，确认教师展示账号挂接到对应系部。',
], [
    '角色编辑页正确回显已绑定角色；保存后再次进入仍保持一致。',
    '菜单权限、流程权限、数据权限分别展示，原子权限归属正确菜单。',
    '组织范围限制对普通账号生效，超级管理员不作为普通权限验收依据。',
])
add_module(doc, '5.2 菜单、权限定义、字典与审计', 'showcase.admin', '管理后台 → 菜单管理 / 权限定义 / 字典管理 / 审计日志', '树形菜单、分页权限定义、教育行业字典和操作审计', [
    '在菜单管理使用树形表格展开教育业务中心，检查父子层级和路由。',
    '在权限定义按菜单权限、流程权限、数据权限筛选并分页查询。',
    '检查部门类型、岗位类别、职级序列、员工类型、在职状态和性别等字典。',
    '执行一次用户编辑或公告发布，再到审计日志按操作者与时间筛选。',
], ['菜单数量与角色授权页一致。', '下拉选项来自字典，数据库与 Redis 缓存一致。', '关键增删改操作生成可检索审计记录。'])

add_heading(doc, '6. 门户与通知公告', 1)
add_module(doc, '6.1 门户工作台', '依次使用 teacher / student / parent', '门户首页 /portal', '我的待办、我的督办、快捷应用、课表、作业、考试、成绩与家校入口', [
    '教师登录后检查待办、课表、会议、监考与教学中心入口。',
    '学生登录后检查个人课表、作业、考试、成绩和请假入口。',
    '家长登录后检查家校通知、学生关联信息和回执入口。',
], ['不同身份看到的卡片和应用符合角色。', '门户数据来自流程中心与教育业务，不出现空白占位或 403。'])
add_module(doc, '6.2 通知公告', 'showcase.admin 发布；其他账号阅读', '管理后台 → 通知公告；门户 → 通知公告', '2 条公告、附件、受众范围、投递记录、用户偏好和渠道模板', [
    '打开公告列表查看草稿/发布状态、有效期、内容类型和可见范围。',
    '新建一条富文本通知，选择组织范围和定时发布时间，保存后预览。',
    '用目标范围账号登录门户查看详情并标记已读。',
    '返回后台查看接收人、投递状态和阅读统计。',
], ['非目标组织用户不可见。', '无有效期时长期有效；定时撤下后门户不再展示。', '附件引用文件中心记录，渠道频率限制与用户偏好生效。'], '短信和企业微信只预留渠道接口；未配置真实凭据时不应执行外部发送。')

add_heading(doc, '7. 流程中心与表单', 1)
add_module(doc, '7.1 流程设计与发布', 'showcase.admin', '管理后台 → 流程中心 / 表单设计器', '10 个流程定义、表单、开始/结束节点、人工/自动/控制节点与连线', [
    '打开已有流程，确认画布加载节点与连线，拖动节点并保存后重新进入。',
    '新增流程时确认自动生成开始节点和结束节点。',
    '配置主表单、节点字段权限、候选人、超时、重试次数与间隔。',
    '校验后发布流程，检查门户发起范围和流程权限。',
], ['保存后节点、坐标、连线和自动节点参数完整回显。', '入口节点和流程分类通过选项选择。', '未通过校验的孤立节点、无出口分支或非法表达式禁止发布。'])
add_module(doc, '7.2 发起与人工审批', 'showcase.teacher / showcase.student 发起；showcase.reviewer 审批', '门户 → 发起流程 / 我的待办', '请假、差旅、采购等实例和任务', [
    '发起人选择流程，填写主表单并上传附件，提交。',
    '审批人进入我的待办查看表单与历史，执行通过、拒绝、退回、转办或加签。',
    '如退回发起人，发起人只能修改并重新提交。',
    '流程结束后在已办和实例详情核对完整轨迹。',
], ['候选人或当前处理人才能办理任务。', '字段权限按节点控制可见、只读和必填。', '每次操作产生审计、通知与流程历史。'])
add_module(doc, '7.3 自动节点与事故恢复', 'showcase.admin 或 WORKFLOW_ADMIN', '管理后台 → 流程事故 / 消息发件箱', '自动节点执行日志、重试、跳过、终止和已解决事故', [
    '选择包含失败自动节点的流程实例，查看事故错误、重试次数、下次重试时间。',
    '执行重试并观察执行日志；仍失败时执行跳过，确认后续人工节点生成。',
    '另选事故执行终止，确认流程不再流转。',
    '查看事故状态、解决方式、解决人和解决时间。',
], ['跳过后事故必须变为已解决，而不是持续 OPEN。', '重试、跳过和终止均满足权限与幂等控制。', '消息失败进入死信并支持重放。'])

add_heading(doc, '8. 教务基础数据', 1)
add_module(doc, '8.1 学年学期、专业、课程与班级', 'showcase.admin / showcase.reviewer', '管理后台 → 教务中心', '3 个学期、6 个专业、16 门课程、16 个科目、6 个行政班', [
    '依次查看学期、专业、课程、科目和行政班列表。',
    '检查当前学期 2026-2027-1、学期日期和启用状态。',
    '打开一个专业和行政班，核对学制、年级、班主任和所属组织。',
], ['上下游引用使用稳定 ID，名称正确回显。', '停用的基础数据不进入新建业务下拉选项。'])
add_module(doc, '8.2 师生、家长和教学任务', 'showcase.admin', '管理后台 → 教师 / 学生 / 家长 / 教学任务 / 教学班成员', '14 名教师、60 名学生、60 名家长、17 个教学任务、300 条成员关系', [
    '查看展示教师和审核人档案，核对员工、部门和任职状态。',
    '查看学生档案、行政班、学籍状态和监护人。',
    '打开教学任务，检查教师、课程、学期、教学班和学生成员。',
    '检查合班公共课教学任务及其来源班级。',
], ['账号、人员档案和业务档案边界清晰。', '教学班成员用于排课、作业、考试和成绩范围。', '合班任务保留各来源行政班，不复制学生档案。'])

add_heading(doc, '9. 走班排课', 1)
add_module(doc, '9.1 资源和规则配置', 'showcase.admin / 排课人员', '管理后台 → 走班排课', '节次、13 间教室、教师禁排、教室不可用、校历例外、规则与约束', [
    '查看周课表的网格/列表切换，确认控件对齐。',
    '检查节次、教室容量与类型、教师每日/每周上限。',
    '查看教师禁排时间、排课偏好、教室不可用时段和节假日。',
    '打开排课策略，检查硬约束与软约束权重。',
], ['禁排、容量、教师冲突、班级冲突和教室冲突作为硬约束。', '单双周、连堂、跨节次和合班任务可表达。'])
add_module(doc, '9.2 生成、调整与发布', 'showcase.admin / 排课人员', '走班排课 → 方案生成 / 冲突检测 / 周课表', '生成任务、候选方案、版本、局部调课和 21 条课表', [
    '启动排课生成任务，观察进度、评分和冲突数量。',
    '预览候选方案，使用教师/班级/教室多维视图检查课表。',
    '对单节课执行拖拽或局部调整，再运行冲突检测。',
    '保存新版本并发布，在教师和学生门户查看结果。',
], ['未解决硬冲突时禁止发布。', '调课保留版本、原因和操作审计。', '考试停课或场地占用能影响冲突判断。'])

add_heading(doc, '10. 教学中心', 1)
add_module(doc, '10.1 教学计划与教案', 'showcase.teacher 编制；showcase.reviewer 审核', '教学中心 → 教学计划 / 教案管理', '教学计划、计划版本、计划项、教案版本与审核记录', [
    '教师按教学任务创建或打开教学计划，维护目标、周次和教学内容。',
    '提交审核后切换审核人处理，记录意见并通过或退回。',
    '教师基于课表创建教案，维护目标、重难点、过程和附件。',
    '生成新版本并比较历史，确认发布后门户可见。',
], ['教学任务是统一入口，不重复创建任务实体。', '主表、审核记录、版本快照职责分离。', 'scheduleEntryId 只读关联，不由教案修改课表。'])
add_module(doc, '10.2 集体备课、课件与教学资源', 'showcase.teacher', '教学中心 → 备课管理 / 课件管理 / 教学资源', '备课成员、评论、材料，课件与版本，文件中心引用', [
    '进入集体备课查看成员、讨论、材料和状态。',
    '上传课件文件并创建版本，绑定到教学任务、教案或备课活动。',
    '预览或下载附件，检查文件权限和引用状态。',
], ['教学班学生成员保持只读，备课成员单独维护。', '文件先上传到文件中心，再进行业务绑定；失败时可补偿。', '历史资源停止双写，兼容映射可追溯。'])
add_module(doc, '10.3 题库、知识点、错题与教研', 'showcase.teacher / showcase.reviewer', '教学中心 → 题库 / 知识点 / 错题 / 教研管理', '题库、题目版本、选项、知识点关联、错题复习、教研组和活动', [
    '创建或查看题库，打开题目、选项、答案、难度和版本。',
    '将题目关联知识点，确认作业和考试仅引用已发布题目版本。',
    '查看由可信评分事件形成的错题，添加复习记录。',
    '查看教研组、成员和教研活动，上传材料并登记成果。',
], ['题目修改生成新版本，不影响历史作业或试卷。', '错题事件具备幂等键；人工录入明确标识来源。', '教研活动形成成员、过程材料和成果闭环。'])

add_heading(doc, '11. 作业中心', 1)
add_module(doc, '11.1 作业发布、提交与批改', 'showcase.teacher / showcase.student', '教学中心 → 作业管理；门户 → 我的作业', '2 份作业及对应提交、批改和发布状态', [
    '教师选择教学任务创建作业，设置题目、截止时间、附件和发布范围。',
    '学生门户打开作业，填写答案或上传附件并提交。',
    '教师查看提交列表，评分、评语并发布成绩。',
    '学生重新进入查看评分和反馈。',
], ['只向教学班成员发布。', '迟交、未交、已批改状态准确。', '发布评分后产生正式事件，供成绩中心和错题本消费。'])

add_heading(doc, '12. 考试中心', 1)
add_module(doc, '12.1 考试计划与资源冲突', 'showcase.admin / 考试管理员', '考试中心 → 考试计划管理', '3 个考试计划、4 个场次、教学任务关联和停课影响', [
    '新建或查看学期考试计划，设置报名/考试日期和适用范围。',
    '添加考试科目与场次，关联教学任务。',
    '运行冲突检查，查看与课表、教室占用和教师时间冲突。',
    '发布考试计划并检查门户可见性。',
], ['同一教师、学生、考场在同一时间不可重复安排。', '考试占用与走班课表共享资源冲突判断。'])
add_module(doc, '12.2 考场、监考与临时调换', 'showcase.admin 排班；showcase.teacher 申请', '考试中心 → 考场管理 / 监考排班与调换；门户 → 我的监考', '3 个考场、监考安排、教师资格和调换记录', [
    '为场次分配考场、容量和考生，检查座位与人数。',
    '按每场所需主监考/副监考数量自动或人工排班。',
    '教师在门户提交不能监考的调换申请并指定候选人或原因。',
    '管理员审核并生成新安排，保留原安排和审批轨迹。',
], ['排班避开教师授课、禁排和其他监考。', '调换前后责任人、原因、时间和审批人可追溯。'])
add_module(doc, '12.3 试卷与考试分析', 'showcase.teacher / showcase.reviewer', '考试中心 → 试卷分析', '试题、逐题得分、知识点和难度统计', [
    '打开考试试卷，检查题目版本、分值和知识点。',
    '查看逐题得分、得分率、区分度和异常题目。',
    '按班级、科目或知识点筛选并导出。',
], ['分析基于已确认考试成绩，候选人稳定映射到学生档案。', '历史试卷保留题目版本快照。'])

add_heading(doc, '13. 成绩中心', 1)
add_module(doc, '13.1 考核方案与成绩册', 'showcase.teacher 编制；showcase.reviewer 审核发布', '成绩中心 → 成绩管理', '考核方案、考核项、成绩册、学生快照、成绩项和课程成绩', [
    '查看教学任务的考核方案，确认平时、作业、考试权重合计为 100%。',
    '生成成绩册成员快照，导入考试成绩并汇总作业评分。',
    '教师提交成绩册，审核人审核后由发布角色发布。',
    '学生门户查看已发布成绩，未发布成绩不可见。',
], ['教师提交、教务审核、独立发布权限可分离。', '成绩册成员快照保留来源、时间和学生关键字段。', '发布后更正必须走更正申请，不直接覆盖历史。'])
add_module(doc, '13.2 成绩分析与变更闭环', 'showcase.reviewer', '成绩中心 → 班级 / 年级 / 学科 / 趋势 / 知识点分析', '成绩更正、变更事故、补考重修、指标快照和多维分析', [
    '分别打开班级、年级、学科、趋势和知识点分析页面。',
    '筛选学期、专业、班级、科目和考试，核对平均分、及格率和分布。',
    '提交成绩更正申请，审核通过后查看新旧值、原因和版本。',
    '查看补考或重修记录及最终处理状态。',
], ['分析口径与已发布成绩一致。', '更正保留原值、现值、申请人、审核人和原因。', '补考重修不破坏原始考试成绩。'], '等级和绩点必须由学校规则集配置，不能在代码中臆造固定算法。')

add_heading(doc, '14. 会议中心', 1)
add_module(doc, '14.1 会议室与会议全周期', 'showcase.teacher 发起；showcase.admin 审批', '会议中心 → 会议室维护 / 会议管理；门户 → 我的会议', '会议室、1 个会议、参会人、材料、纪要和行动项', [
    '管理员维护会议室容量、设备、位置和可用状态。',
    '教师创建会议，选择会议室、时间、参会人并提交审批。',
    '管理员批准后，参会人门户收到邀请并确认。',
    '会后上传材料、发布纪要并分派行动项，跟踪完成状态。',
], ['会议室时间冲突时禁止提交或批准。', '会议审批权限允许管理员处理，不出现无权访问。', '线上会议链接作为扩展字段，可对接第三方服务。'])

add_heading(doc, '15. 请假与教室预约', 1)
add_module(doc, '15.1 学生请假', 'showcase.student 发起；showcase.teacher 审批', '门户 → 请假；管理后台 → 请假管理', '请假申请、审批状态和返校销假字段', [
    '学生填写请假类型、起止时间、原因和附件并提交。',
    '班主任查看申请，审批通过或退回。',
    '学生查看状态；返校后执行销假。',
], ['时间校验正确，跨课次请假影响考勤提示。', '审批历史与附件可追溯。'])
add_module(doc, '15.2 教室预约', 'showcase.teacher', '门户 → 教室预约', '教室资源、预约记录、不可用时段和冲突', [
    '选择日期、节次、人数和教室类型查询可用教室。',
    '提交预约并查看审批状态。',
    '尝试选择已有课表或考试占用的教室时间。',
], ['可用查询综合课表、考试、会议和教室不可用时段。', '冲突预约被阻止，成功预约可取消并释放资源。'])

add_heading(doc, '16. 家校协同', 1)
add_module(doc, '16.1 班级通知与家长回执', 'showcase.teacher 发布；showcase.parent 回执', '门户 → 班主任工作台 / 班级通知；家长门户 → 家校协同', '家长绑定、班级通知、接收人、回执和家庭通知', [
    '班主任创建班级通知，选择班级、截止时间并要求回执。',
    '家长登录查看与绑定学生相关的通知并回执。',
    '班主任查看已读、未读、已回执和未回执名单。',
], ['家长只能看到绑定学生范围内的数据。', '通知接收人生成稳定，重复发布不会重复接收记录。', '回执率和未回执名单可用于提醒。'])

add_heading(doc, '17. 教学督导与数据中心', 1)
add_module(doc, '17.1 教学督导', 'showcase.reviewer', '管理后台 → 教学督导；门户 → 我的督导', '督导模板、计划、任务、听课记录、问题、整改与复核', [
    '创建督导计划，选择模板、范围和督导员。',
    '督导员按任务进入课堂评价，填写指标、意见和附件。',
    '对发现的问题发起整改，责任人提交整改材料。',
    '督导人复核并关闭问题。',
], ['计划—任务—记录—问题—整改—复核形成闭环。', '普通用户只能看到本人相关任务。'])
add_module(doc, '17.2 数据指标、质量和报表', 'showcase.admin / showcase.reviewer', '管理后台 → 数据中心', '成绩事实、指标快照、2 个质量问题和报表任务', [
    '查看指标目录和当前快照，按学期、组织、专业筛选。',
    '进入质量问题查看规则、严重度、责任人和处理状态。',
    '处理问题并复核关闭，观察指标或快照更新。',
    '创建报表任务并查看执行状态和产物。',
], ['指标口径可追溯到定义与数据周期。', '质量问题支持发现、分派、修复、复核和关闭。'])

add_heading(doc, '18. AI、知识库与集成中心', 1)
add_module(doc, '18.1 AI 模型与知识库', 'showcase.admin', '管理后台 → 模型管理 / AI 接口 / 知识库', '模型配置元数据、知识库、1 份文档及向量化链路', [
    '在模型管理配置供应商、模型和密钥引用，执行连通性测试。',
    '创建知识库，上传文档，观察解析、分块和向量化状态。',
    '在知识助手中提问并核对引用来源。',
], ['应用启动不强制要求 yml 中存在 DeepSeek Key。', '密钥从模型管理的加密配置读取，接口不返回明文。', 'RAG 回答带文档引用；模型不可用时给出明确状态。'], '知识库技术路线由 Embedding + 向量检索 + RAG 实现；不要在展示数据中写入真实 API Key。')
add_module(doc, '18.2 集成中心', 'showcase.admin', '管理后台 → 集成中心', '2 个连接器、同步任务、运行记录、条目错误和死信', [
    '查看连接器类型、启用状态和最近健康检查。',
    '打开同步任务与运行详情，检查成功数、失败数和错误明细。',
    '查看死信并执行重试或标记解决。',
], ['列表和运行详情即使无真实凭据也可展示。', '凭据使用 int_credential 加密保存，不由 SQL 脚本生成。', '重试有次数限制、幂等键和审计。'])

add_heading(doc, '19. 验收清单与问题定位', 1)
add_heading(doc, '19.1 最小验收清单', 2)
add_table(doc, ['编号', '验收项', '通过标准'], [
    ('A01', '登录与导航', '五个账号均可登录；菜单符合角色；无 401/403 循环'),
    ('A02', '基础数据', '组织、师生、课程、班级下拉均正确显示中文名称'),
    ('A03', '排课', '多维课表可见；硬冲突被阻止；发布后门户可见'),
    ('A04', '教学', '计划、教案、备课、课件、题库、知识点、错题、教研页面可用'),
    ('A05', '作业', '教师发布—学生提交—教师批改—学生查看闭环'),
    ('A06', '考试', '计划、场次、考场、监考、调换、试卷分析闭环'),
    ('A07', '成绩', '方案—汇总—审核—发布—更正—分析闭环'),
    ('A08', '流程', '发起—审批—退回/重提—结束；事故恢复状态正确'),
    ('A09', '消息', '范围、定时发布、有效期、阅读与投递统计正确'),
    ('A10', '会议与资源', '会议室、教室冲突控制和审批可用'),
    ('A11', '家校与督导', '家长隔离、回执统计、督导整改闭环'),
    ('A12', '数据与集成', '指标、质量、报表、同步错误和死信可查询'),
], [1.5, 4.3, 10.2])

add_heading(doc, '19.2 常见问题定位', 2)
add_table(doc, ['现象', '优先检查', '处理建议'], [
    ('登录后一闪返回登录页', 'Token、/portal/bootstrap、账号状态、token_version', '清除旧 Token，检查鉴权过滤器与账号状态'),
    ('菜单比授权页少', '菜单启用状态、角色菜单、权限所属菜单', '修复菜单绑定后重新登录刷新导航'),
    ('下拉显示 ID 或无数据', '字典缓存、名称映射接口、关联记录', '同步数据库与 Redis 字典，禁止前端写死选项'),
    ('保存权限时请求头过大', 'Token 权限声明体积', 'Token 只保存身份与版本，权限服务端查询并缓存'),
    ('流程画布缺节点或连线', '节点/边接口、LOB 映射、flowId', '检查两类数据均加载，避免把大字段映射为 OID'),
    ('事故跳过后仍 OPEN', 'resolution、resolved_by、resolved_at', '跳过事务内同时完成事故关闭与后继推进'),
    ('启动提示 Flyway 校验失败', 'flyway_schema_history 与迁移文件校验和', '新增补偿迁移；不要修改已执行迁移或盲目 repair'),
    ('AI 启动要求 API Key', '自动配置是否被强制加载', '从模型管理延迟创建模型，未配置时返回 unavailable'),
    ('附件不可预览', 'MinIO、managed_file 状态、业务绑定', '先检查对象存储，再检查两阶段文件绑定'),
], [4.2, 5.2, 6.6])

add_heading(doc, '19.3 演示后的数据处理', 2)
add_bullets(doc, [
    '本脚本设计为幂等追加与更新，不提供 destructive 清理命令。',
    '需要恢复演示基线时，优先重建独立 ChronosEducation 演示数据库并重新执行 Flyway 与本脚本。',
    '不要在共享库通过模糊条件批量删除 showcase_seed 数据；流程、消息和附件可能存在引用。',
    '正式上线前删除展示账号，轮换所有密钥，清理模拟通知、流程和外部连接器。',
])

add_heading(doc, '附录 A：演示口令与操作记录表', 1)
add_note(doc, '统一密码', 'ChronosExam@2026（仅本地演示环境）', 'E2F0D9')
add_table(doc, ['日期', '执行人', '场景', '结果', '问题单'], [('____', '____', '____', '通过 / 不通过', '____') for _ in range(8)], [2.5, 2.5, 5, 3, 3])

add_heading(doc, '附录 B：交付物', 1)
add_bullets(doc, [
    '完整展示数据脚本：Chronos/scripts/education-complete-showcase-data.sql。',
    '本操作手册：docs/Chronos教育行业全业务展示与操作手册.docx。',
    '脚本已在 ChronosEducation 实际数据库成功执行并完成二次幂等复跑。',
])

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(OUTPUT)
